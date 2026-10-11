"""Sales & operations report generator for small businesses.

Reads three CSV files (sales, expenses, inventory) and builds one Excel
workbook with a summary dashboard, sales by category, product ranking,
expense breakdown and an inventory status sheet with reorder alerts.

Every figure in the report sheets is an Excel formula that points to the
data sheets, so the workbook keeps working if someone edits the data.

Usage:
    python report_generator.py
    python report_generator.py --sales my_sales.csv --expenses my_expenses.csv \
        --inventory my_inventory.csv --output my_report.xlsx
"""
import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent
FONT = "Arial"
NAVY = "1F3A5F"
LIGHT = "EAF0F7"
GREY = "F5F5F5"

REQUIRED_COLUMNS = {
    "sales": ["date", "product_id", "product", "category", "units", "unit_price"],
    "expenses": ["date", "category", "description", "amount"],
    "inventory": ["product_id", "product", "category", "unit_cost", "unit_price",
                  "stock_on_hand", "reorder_level", "supplier"],
}

MONEY = '$#,##0.00'
MONEY0 = '$#,##0'
PCT = '0.0%'
MONTH_FMT = 'mmm yyyy'
DATE_FMT = 'yyyy-mm-dd'


class DataError(Exception):
    """Raised when an input file is missing, malformed or inconsistent."""


# --------------------------------------------------------------------------
# Loading and validation
# --------------------------------------------------------------------------
def load_csv(path, kind):
    path = Path(path)
    if not path.exists():
        raise DataError(f"File not found: {path}")
    try:
        df = pd.read_csv(path)
    except Exception as exc:  # malformed CSV
        raise DataError(f"Could not read {path.name}: {exc}") from exc

    missing = [c for c in REQUIRED_COLUMNS[kind] if c not in df.columns]
    if missing:
        raise DataError(f"{path.name} is missing columns: {', '.join(missing)}")
    if df.empty:
        raise DataError(f"{path.name} has no rows")
    return df[REQUIRED_COLUMNS[kind]].copy()


def to_numeric(df, columns, filename):
    for col in columns:
        converted = pd.to_numeric(df[col], errors="coerce")
        if converted.isna().any():
            bad = int(converted.isna().sum())
            raise DataError(f"{filename}: column '{col}' has {bad} non-numeric value(s)")
        df[col] = converted


def to_dates(df, filename):
    parsed = pd.to_datetime(df["date"], errors="coerce")
    if parsed.isna().any():
        raise DataError(f"{filename}: column 'date' has {int(parsed.isna().sum())} invalid date(s)")
    df["date"] = parsed.dt.date


def load_all(sales_path, expenses_path, inventory_path):
    sales = load_csv(sales_path, "sales")
    expenses = load_csv(expenses_path, "expenses")
    inventory = load_csv(inventory_path, "inventory")

    to_dates(sales, Path(sales_path).name)
    to_dates(expenses, Path(expenses_path).name)
    to_numeric(sales, ["units", "unit_price"], Path(sales_path).name)
    to_numeric(expenses, ["amount"], Path(expenses_path).name)
    to_numeric(inventory, ["unit_cost", "unit_price", "stock_on_hand", "reorder_level"],
               Path(inventory_path).name)

    unknown = sorted(set(sales["product_id"]) - set(inventory["product_id"]))
    if unknown:
        raise DataError("Sales contain product IDs that are not in the inventory file: "
                        + ", ".join(map(str, unknown[:10])))
    if inventory["product_id"].duplicated().any():
        raise DataError("Inventory has duplicated product IDs")

    sales = sales.sort_values(["date", "product_id"]).reset_index(drop=True)
    expenses = expenses.sort_values("date").reset_index(drop=True)
    return sales, expenses, inventory


# --------------------------------------------------------------------------
# Styling helpers
# --------------------------------------------------------------------------
def style_header(ws, row, first_col, last_col):
    fill = PatternFill("solid", fgColor=NAVY)
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_body(ws, first_row, last_row, first_col, last_col, zebra=True):
    thin = Side(style="thin", color="D9D9D9")
    for r in range(first_row, last_row + 1):
        for c in range(first_col, last_col + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font.name != FONT:
                cell.font = Font(name=FONT, size=10)
            cell.border = Border(bottom=thin)
            if zebra and (r - first_row) % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=GREY)


def style_total_row(ws, row, first_col, last_col):
    top = Side(style="medium", color=NAVY)
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.border = Border(top=top)


def set_widths(ws, widths):
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=FONT, size=9, italic=True, color="666666")


def first_of_month(d):
    return d.replace(day=1)


def rng(sheet, col, last_row):
    """Absolute range such as Data_Sales!$H$2:$H$1473."""
    return f"{sheet}!${col}$2:${col}${last_row}"


def finish_chart(chart, title_text, width=17, height=8, y_title=None, x_title=None):
    chart.title = title_text
    chart.width = width
    chart.height = height
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    if y_title:
        chart.y_axis.title = y_title
    if x_title:
        chart.x_axis.title = x_title
    return chart


# --------------------------------------------------------------------------
# Data sheets (raw data + helper formulas)
# --------------------------------------------------------------------------
def build_data_sales(wb, sales, n_inv):
    ws = wb.create_sheet("Data_Sales")
    headers = ["Date", "Month", "Product ID", "Product", "Category", "Units",
               "Unit Price", "Revenue", "Unit Cost", "COGS", "Gross Profit"]
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    inv_id = rng("Inventory", "A", n_inv + 1)
    inv_cost = rng("Inventory", "D", n_inv + 1)
    for i, row in enumerate(sales.itertuples(index=False), start=2):
        ws.append([
            row.date,
            f"=DATE(YEAR(A{i}),MONTH(A{i}),1)",
            row.product_id, row.product, row.category,
            row.units, row.unit_price,
            f"=F{i}*G{i}",
            f"=INDEX({inv_cost},MATCH(C{i},{inv_id},0))",
            f"=F{i}*I{i}",
            f"=H{i}-J{i}",
        ])
    last = len(sales) + 1
    for r in range(2, last + 1):
        ws.cell(r, 1).number_format = DATE_FMT
        ws.cell(r, 2).number_format = MONTH_FMT
        for c in (7, 8, 9, 10, 11):
            ws.cell(r, c).number_format = MONEY
        for c in range(1, 12):
            ws.cell(r, c).font = Font(name=FONT, size=10)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{last}"
    set_widths(ws, {"A": 12, "B": 11, "C": 11, "D": 22, "E": 13, "F": 8,
                    "G": 11, "H": 11, "I": 11, "J": 11, "K": 12})
    return last


def build_data_expenses(wb, expenses):
    ws = wb.create_sheet("Data_Expenses")
    headers = ["Date", "Month", "Category", "Description", "Amount"]
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    for i, row in enumerate(expenses.itertuples(index=False), start=2):
        ws.append([row.date, f"=DATE(YEAR(A{i}),MONTH(A{i}),1)",
                   row.category, row.description, row.amount])
    last = len(expenses) + 1
    for r in range(2, last + 1):
        ws.cell(r, 1).number_format = DATE_FMT
        ws.cell(r, 2).number_format = MONTH_FMT
        ws.cell(r, 5).number_format = MONEY
        for c in range(1, 6):
            ws.cell(r, c).font = Font(name=FONT, size=10)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{last}"
    set_widths(ws, {"A": 12, "B": 11, "C": 14, "D": 30, "E": 12})
    return last


# --------------------------------------------------------------------------
# Inventory sheet
# --------------------------------------------------------------------------
def build_inventory(wb, inventory, last_sales):
    ws = wb.create_sheet("Inventory")
    headers = ["Product ID", "Product", "Category", "Unit Cost", "Unit Price",
               "Stock On Hand", "Reorder Level", "Status", "Stock Value",
               "Margin %", "Units Sold (period)", "Avg Daily Sales", "Days of Stock Left"]
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    sales_units = rng("Data_Sales", "F", last_sales)
    sales_ids = rng("Data_Sales", "C", last_sales)
    blue = Font(name=FONT, size=10, color="0000FF")
    for i, row in enumerate(inventory.itertuples(index=False), start=2):
        ws.append([
            row.product_id, row.product, row.category, row.unit_cost, row.unit_price,
            row.stock_on_hand, row.reorder_level,
            f'=IF(F{i}<=G{i},"REORDER","OK")',
            f"=F{i}*D{i}",
            f"=IF(E{i}=0,0,(E{i}-D{i})/E{i})",
            f"=SUMIFS({sales_units},{sales_ids},A{i})",
            f"=K{i}/Summary!$B$5",
            f'=IF(L{i}=0,"",ROUND(F{i}/L{i},1))',
        ])
    last = len(inventory) + 1
    for r in range(2, last + 1):
        for c in range(1, 14):
            ws.cell(r, c).font = blue if c <= 7 and c != 8 else Font(name=FONT, size=10)
        ws.cell(r, 4).number_format = MONEY
        ws.cell(r, 5).number_format = MONEY
        ws.cell(r, 9).number_format = MONEY
        ws.cell(r, 10).number_format = PCT
        ws.cell(r, 12).number_format = '0.0'
        ws.cell(r, 13).number_format = '0.0'
        ws.cell(r, 8).alignment = Alignment(horizontal="center")
    total = last + 1
    ws.cell(total, 1, "Total")
    ws.cell(total, 9, f"=SUM(I2:I{last})")
    ws.cell(total, 9).number_format = MONEY
    style_total_row(ws, total, 1, 13)
    ws.cell(total, 9).number_format = MONEY

    red_fill = PatternFill("solid", bgColor="F8D7DA", fgColor="F8D7DA")
    ws.conditional_formatting.add(
        f"A2:M{last}",
        FormulaRule(formula=['$H2="REORDER"'], fill=red_fill,
                    font=Font(name=FONT, bold=True, color="9C0006")))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{last}"
    set_widths(ws, {"A": 11, "B": 22, "C": 12, "D": 11, "E": 11, "F": 12, "G": 12,
                    "H": 11, "I": 13, "J": 10, "K": 14, "L": 12, "M": 14})
    ws.row_dimensions[1].height = 32
    return last


# --------------------------------------------------------------------------
# Report sheets
# --------------------------------------------------------------------------
def build_summary(wb, months, last_sales, last_exp, last_inv, store_name):
    ws = wb["Summary"]
    title(ws, f"Sales & Operations Report - {store_name}",
          "All figures are formulas linked to the Data_* and Inventory sheets.")

    sales_rev = rng("Data_Sales", "H", last_sales)
    sales_cogs = rng("Data_Sales", "J", last_sales)
    sales_date = rng("Data_Sales", "A", last_sales)
    sales_month = rng("Data_Sales", "B", last_sales)
    exp_amt = rng("Data_Expenses", "E", last_exp)
    exp_month = rng("Data_Expenses", "B", last_exp)
    inv_status = rng("Inventory", "H", last_inv)

    rows = [
        (3, "Period start", f"=MIN({sales_date})", DATE_FMT),
        (4, "Period end", f"=MAX({sales_date})", DATE_FMT),
        (5, "Days in period", "=B4-B3+1", '0'),
        (7, "Total Revenue", f"=SUM({sales_rev})", MONEY0),
        (8, "Cost of Goods Sold", f"=SUM({sales_cogs})", MONEY0),
        (9, "Gross Profit", "=B7-B8", MONEY0),
        (10, "Gross Margin %", "=IF(B7=0,0,B9/B7)", PCT),
        (11, "Operating Expenses", f"=SUM({exp_amt})", MONEY0),
        (12, "Net Profit", "=B9-B11", MONEY0),
        (13, "Net Margin %", "=IF(B7=0,0,B12/B7)", PCT),
        (14, "Products to Reorder", f'=COUNTIF({inv_status},"REORDER")', '0'),
    ]
    for r, label, formula, fmt in rows:
        ws.cell(r, 1, label).font = Font(name=FONT, size=10, bold=r >= 7)
        c = ws.cell(r, 2, formula)
        c.number_format = fmt
        c.font = Font(name=FONT, size=11 if r >= 7 else 10, bold=r >= 7)
        c.alignment = Alignment(horizontal="right")
    for r in (7, 9, 12):
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(r, 2).fill = PatternFill("solid", fgColor=LIGHT)
    ws["B14"].font = Font(name=FONT, size=11, bold=True, color="9C0006")

    # Monthly table
    start = 18
    headers = ["Month", "Revenue", "COGS", "Gross Profit", "Operating Expenses",
               "Net Profit", "Net Margin %"]
    ws.cell(start - 1, 1, "Monthly performance").font = Font(name=FONT, size=12, bold=True, color=NAVY)
    for j, h in enumerate(headers, start=1):
        ws.cell(start, j, h)
    style_header(ws, start, 1, len(headers))
    for i, m in enumerate(months, start=start + 1):
        ws.cell(i, 1, m).number_format = MONTH_FMT
        ws.cell(i, 2, f"=SUMIFS({sales_rev},{sales_month},$A{i})")
        ws.cell(i, 3, f"=SUMIFS({sales_cogs},{sales_month},$A{i})")
        ws.cell(i, 4, f"=B{i}-C{i}")
        ws.cell(i, 5, f"=SUMIFS({exp_amt},{exp_month},$A{i})")
        ws.cell(i, 6, f"=D{i}-E{i}")
        ws.cell(i, 7, f"=IF(B{i}=0,0,F{i}/B{i})")
        for c in range(2, 7):
            ws.cell(i, c).number_format = MONEY0
        ws.cell(i, 7).number_format = PCT
    last_m = start + len(months)
    style_body(ws, start + 1, last_m, 1, 7)
    total = last_m + 1
    ws.cell(total, 1, "Total")
    for c in range(2, 7):
        col = get_column_letter(c)
        ws.cell(total, c, f"=SUM({col}{start + 1}:{col}{last_m})").number_format = MONEY0
    ws.cell(total, 7, f"=IF(B{total}=0,0,F{total}/B{total})").number_format = PCT
    style_total_row(ws, total, 1, 7)

    notes_row = total + 2
    notes = [
        "Notes",
        "- Gross Profit = Revenue - COGS. COGS uses the unit cost from the Inventory sheet.",
        "- Net Profit = Gross Profit - Operating Expenses (rent, payroll, utilities, etc.).",
        "- The sample data included in this repository is synthetic.",
    ]
    for k, text in enumerate(notes):
        cell = ws.cell(notes_row + k, 1, text)
        cell.font = Font(name=FONT, size=9, bold=(k == 0), color="555555")

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    for col in (2, 5, 6):  # Revenue, Operating Expenses, Net Profit
        chart.add_data(Reference(ws, min_col=col, min_row=start, max_row=last_m), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=start + 1, max_row=last_m))
    finish_chart(chart, "Revenue, expenses and net profit by month", width=17, height=7.4, y_title="USD")
    chart.x_axis.number_format = MONTH_FMT
    chart.y_axis.number_format = '$#,##0'
    ws.add_chart(chart, "D3")

    set_widths(ws, {"A": 24, "B": 15, "C": 14, "D": 14, "E": 20, "F": 14, "G": 14})
    ws.sheet_view.showGridLines = False
    return ws


def build_categories(wb, categories, months, last_sales):
    ws = wb.create_sheet("Sales_by_Category")
    title(ws, "Sales by Category")
    sales_rev = rng("Data_Sales", "H", last_sales)
    sales_gp = rng("Data_Sales", "K", last_sales)
    sales_cat = rng("Data_Sales", "E", last_sales)
    sales_month = rng("Data_Sales", "B", last_sales)

    ws.cell(3, 1, "Revenue by category and month").font = Font(name=FONT, size=12, bold=True, color=NAVY)
    header_row = 4
    ws.cell(header_row, 1, "Category")
    for j, m in enumerate(months, start=2):
        ws.cell(header_row, j, m).number_format = MONTH_FMT
    tot_col = len(months) + 2
    ws.cell(header_row, tot_col, "Total")
    ws.cell(header_row, tot_col + 1, "% of Revenue")
    style_header(ws, header_row, 1, tot_col + 1)
    for j in range(2, len(months) + 2):
        ws.cell(header_row, j).number_format = MONTH_FMT

    first = header_row + 1
    last = header_row + len(categories)
    for i, cat in enumerate(categories, start=first):
        ws.cell(i, 1, cat)
        for j in range(2, len(months) + 2):
            col = get_column_letter(j)
            ws.cell(i, j, f"=SUMIFS({sales_rev},{sales_cat},$A{i},{sales_month},{col}${header_row})").number_format = MONEY0
        first_col = get_column_letter(2)
        last_col = get_column_letter(len(months) + 1)
        ws.cell(i, tot_col, f"=SUM({first_col}{i}:{last_col}{i})").number_format = MONEY0
        tc = get_column_letter(tot_col)
        ws.cell(i, tot_col + 1, f"=IF({tc}${last + 1}=0,0,{tc}{i}/{tc}${last + 1})").number_format = PCT
    style_body(ws, first, last, 1, tot_col + 1)
    total = last + 1
    ws.cell(total, 1, "Total")
    for j in range(2, tot_col + 1):
        col = get_column_letter(j)
        ws.cell(total, j, f"=SUM({col}{first}:{col}{last})").number_format = MONEY0
    ws.cell(total, tot_col + 1, f"=SUM({get_column_letter(tot_col + 1)}{first}:{get_column_letter(tot_col + 1)}{last})").number_format = PCT
    style_total_row(ws, total, 1, tot_col + 1)

    # Margin table
    m_title = total + 3
    ws.cell(m_title - 1, 1, "Profitability by category").font = Font(name=FONT, size=12, bold=True, color=NAVY)
    heads = ["Category", "Revenue", "Gross Profit", "Gross Margin %"]
    for j, h in enumerate(heads, start=1):
        ws.cell(m_title, j, h)
    style_header(ws, m_title, 1, 4)
    for k, cat in enumerate(categories, start=m_title + 1):
        ws.cell(k, 1, cat)
        ws.cell(k, 2, f"=SUMIFS({sales_rev},{sales_cat},$A{k})").number_format = MONEY0
        ws.cell(k, 3, f"=SUMIFS({sales_gp},{sales_cat},$A{k})").number_format = MONEY0
        ws.cell(k, 4, f"=IF(B{k}=0,0,C{k}/B{k})").number_format = PCT
    m_last = m_title + len(categories)
    style_body(ws, m_title + 1, m_last, 1, 4)

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.add_data(Reference(ws, min_col=2, max_col=len(months) + 1, min_row=header_row, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    finish_chart(chart, "Revenue by category", width=17, height=8, y_title="USD")
    chart.y_axis.number_format = '$#,##0'
    ws.add_chart(chart, "H3")

    set_widths(ws, {"A": 16, "B": 14, "C": 14, "D": 16, "E": 14, "F": 14})
    ws.sheet_view.showGridLines = False


def build_products(wb, products_sorted, last_sales):
    ws = wb.create_sheet("Product_Ranking")
    title(ws, "Product Ranking",
          "Products are listed by revenue at the time the report was generated.")
    sales_units = rng("Data_Sales", "F", last_sales)
    sales_rev = rng("Data_Sales", "H", last_sales)
    sales_gp = rng("Data_Sales", "K", last_sales)
    sales_id = rng("Data_Sales", "C", last_sales)

    headers = ["Rank", "Product ID", "Product", "Category", "Units Sold",
               "Revenue", "Gross Profit", "Gross Margin %", "% of Revenue"]
    header_row = 4
    for j, h in enumerate(headers, start=1):
        ws.cell(header_row, j, h)
    style_header(ws, header_row, 1, len(headers))
    first = header_row + 1
    last = header_row + len(products_sorted)
    total = last + 1
    for i, (pid, name, cat) in enumerate(products_sorted, start=first):
        ws.cell(i, 1, f"=RANK(F{i},$F${first}:$F${last})")
        ws.cell(i, 2, pid)
        ws.cell(i, 3, name)
        ws.cell(i, 4, cat)
        ws.cell(i, 5, f"=SUMIFS({sales_units},{sales_id},$B{i})").number_format = '#,##0'
        ws.cell(i, 6, f"=SUMIFS({sales_rev},{sales_id},$B{i})").number_format = MONEY0
        ws.cell(i, 7, f"=SUMIFS({sales_gp},{sales_id},$B{i})").number_format = MONEY0
        ws.cell(i, 8, f"=IF(F{i}=0,0,G{i}/F{i})").number_format = PCT
        ws.cell(i, 9, f"=IF(F${total}=0,0,F{i}/F${total})").number_format = PCT
        ws.cell(i, 1).alignment = Alignment(horizontal="center")
    style_body(ws, first, last, 1, len(headers))
    ws.cell(total, 1, "Total")
    ws.cell(total, 5, f"=SUM(E{first}:E{last})").number_format = '#,##0'
    ws.cell(total, 6, f"=SUM(F{first}:F{last})").number_format = MONEY0
    ws.cell(total, 7, f"=SUM(G{first}:G{last})").number_format = MONEY0
    ws.cell(total, 8, f"=IF(F{total}=0,0,G{total}/F{total})").number_format = PCT
    ws.cell(total, 9, f"=SUM(I{first}:I{last})").number_format = PCT
    style_total_row(ws, total, 1, len(headers))

    top_n = min(10, len(products_sorted))
    chart = BarChart()
    chart.type = "bar"
    chart.add_data(Reference(ws, min_col=6, min_row=header_row, max_row=header_row + top_n), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=3, min_row=first, max_row=first + top_n - 1))
    finish_chart(chart, f"Top {top_n} products by revenue", width=17, height=9)
    chart.legend = None
    chart.x_axis.scaling.orientation = "maxMin"
    chart.y_axis.number_format = '$#,##0'
    ws.add_chart(chart, "K3")

    ws.freeze_panes = f"A{first}"
    set_widths(ws, {"A": 7, "B": 11, "C": 22, "D": 12, "E": 11, "F": 12, "G": 13, "H": 14, "I": 13})
    ws.sheet_view.showGridLines = False


def build_expenses(wb, expense_cats, months, last_exp):
    ws = wb.create_sheet("Expense_Breakdown")
    title(ws, "Expense Breakdown",
          "Categories are listed by total spend at the time the report was generated.")
    amt = rng("Data_Expenses", "E", last_exp)
    cat = rng("Data_Expenses", "C", last_exp)
    month = rng("Data_Expenses", "B", last_exp)

    header_row = 4
    ws.cell(header_row, 1, "Category")
    for j, m in enumerate(months, start=2):
        ws.cell(header_row, j, m).number_format = MONTH_FMT
    tot_col = len(months) + 2
    ws.cell(header_row, tot_col, "Total")
    ws.cell(header_row, tot_col + 1, "% of Expenses")
    style_header(ws, header_row, 1, tot_col + 1)
    for j in range(2, len(months) + 2):
        ws.cell(header_row, j).number_format = MONTH_FMT

    first = header_row + 1
    last = header_row + len(expense_cats)
    total = last + 1
    tc = get_column_letter(tot_col)
    for i, name in enumerate(expense_cats, start=first):
        ws.cell(i, 1, name)
        for j in range(2, len(months) + 2):
            col = get_column_letter(j)
            ws.cell(i, j, f"=SUMIFS({amt},{cat},$A{i},{month},{col}${header_row})").number_format = MONEY0
        ws.cell(i, tot_col, f"=SUM(B{i}:{get_column_letter(len(months) + 1)}{i})").number_format = MONEY0
        ws.cell(i, tot_col + 1, f"=IF({tc}${total}=0,0,{tc}{i}/{tc}${total})").number_format = PCT
    style_body(ws, first, last, 1, tot_col + 1)
    ws.cell(total, 1, "Total")
    for j in range(2, tot_col + 1):
        col = get_column_letter(j)
        ws.cell(total, j, f"=SUM({col}{first}:{col}{last})").number_format = MONEY0
    ws.cell(total, tot_col + 1, f"=SUM({get_column_letter(tot_col + 1)}{first}:{get_column_letter(tot_col + 1)}{last})").number_format = PCT
    style_total_row(ws, total, 1, tot_col + 1)

    chart = BarChart()
    chart.type = "bar"
    chart.add_data(Reference(ws, min_col=tot_col, min_row=header_row, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    finish_chart(chart, "Total expenses by category", width=17, height=8)
    chart.legend = None
    chart.x_axis.scaling.orientation = "maxMin"
    chart.y_axis.number_format = '$#,##0'
    ws.add_chart(chart, "H3")

    set_widths(ws, {"A": 16, "B": 13, "C": 13, "D": 13, "E": 13, "F": 14})
    ws.sheet_view.showGridLines = False


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def build_report(sales, expenses, inventory, output, store_name):
    months = sorted({first_of_month(d) for d in sales["date"]} |
                    {first_of_month(d) for d in expenses["date"]})

    sales = sales.assign(revenue=sales["units"] * sales["unit_price"])
    products = (sales.groupby(["product_id", "product", "category"])["revenue"].sum()
                .sort_values(ascending=False).reset_index())
    products_sorted = list(products[["product_id", "product", "category"]].itertuples(index=False, name=None))
    # Products in inventory that never sold still show up (at the bottom).
    sold_ids = {p[0] for p in products_sorted}
    for r in inventory.itertuples(index=False):
        if r.product_id not in sold_ids:
            products_sorted.append((r.product_id, r.product, r.category))
    categories = list(sales.groupby("category")["revenue"].sum().sort_values(ascending=False).index)
    expense_cats = list(expenses.groupby("category")["amount"].sum().sort_values(ascending=False).index)

    wb = Workbook()
    wb.active.title = "Summary"
    n_inv = len(inventory)

    last_sales = len(sales) + 1
    last_exp = len(expenses) + 1
    last_inv = n_inv + 1

    build_summary(wb, months, last_sales, last_exp, last_inv, store_name)
    build_categories(wb, categories, months, last_sales)
    build_products(wb, products_sorted, last_sales)
    build_expenses(wb, expense_cats, months, last_exp)
    build_inventory(wb, inventory, last_sales)
    build_data_sales(wb, sales.drop(columns="revenue"), n_inv)
    build_data_expenses(wb, expenses)

    # Order: reports first, data last.
    order = ["Summary", "Sales_by_Category", "Product_Ranking", "Expense_Breakdown",
             "Inventory", "Data_Sales", "Data_Expenses"]
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def parse_args():
    p = argparse.ArgumentParser(description="Generate an Excel sales & operations report from CSV files.")
    p.add_argument("--sales", default=BASE_DIR / "data" / "sample_sales.csv")
    p.add_argument("--expenses", default=BASE_DIR / "data" / "sample_expenses.csv")
    p.add_argument("--inventory", default=BASE_DIR / "data" / "sample_inventory.csv")
    p.add_argument("--output", default=BASE_DIR / "output" / "sales_report.xlsx")
    p.add_argument("--store-name", default="Sample Grocery Store")
    return p.parse_args()


def main():
    args = parse_args()
    try:
        sales, expenses, inventory = load_all(args.sales, args.expenses, args.inventory)
    except DataError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    out = build_report(sales, expenses, inventory, args.output, args.store_name)
    print(f"Report created: {out}")
    print(f"  Sales rows: {len(sales)} | Expense rows: {len(expenses)} | Products: {len(inventory)}")
    print("Note: formulas calculate when the file is opened in Excel or LibreOffice.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
